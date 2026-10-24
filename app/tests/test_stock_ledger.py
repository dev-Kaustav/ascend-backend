"""Stock ledger: running balances, the opening/closing summary, and the Excel export."""
from datetime import datetime, timezone
from io import BytesIO

import openpyxl

from app.core.security import create_access_token, get_password_hash
from app.models import Brand, InventoryTransaction, Order, Retailer, SKU, User, Warehouse
from app.models.enums import EmployeeRole, OrderStatus, TransactionType
from app.schemas.admin import InventoryReceiptCreate, InventoryReceiptItem
from app.services.admin import add_inventory_receipt
from app.services.stock_ledger import get_stock_ledger


def _seed(db):
    brand = Brand(name="Ledger Brand")
    warehouse = Warehouse(name="Ledger WH", location="Delhi", state="Delhi")
    retailer = Retailer(name="Ledger Retailer", state="Delhi")
    db.add_all([brand, warehouse, retailer])
    db.flush()
    sku = SKU(name="Ledger Chips", code="LDG-1", brand_id=brand.id)
    db.add(sku)
    db.commit()
    return brand, warehouse, retailer, sku


def _move(db, sku, warehouse, kind, quantity, when, order_id=None):
    db.add(InventoryTransaction(
        sku_id=sku.id, warehouse_id=warehouse.id, transaction_type=kind, quantity=quantity,
        created_at=when, order_id=order_id,
    ))
    db.commit()


def test_receipt_records_its_order_and_supplier(db):
    brand, warehouse, _, sku = _seed(db)
    order = add_inventory_receipt(db, InventoryReceiptCreate(
        brand_id=brand.id, warehouse_id=warehouse.id,
        items=[InventoryReceiptItem(sku_id=sku.id, quantity=40, mfg_date=None, expiry_date=None)],
    ))

    movement = get_stock_ledger(db)["movements"][0]

    assert movement["order_id"] == order.id
    assert movement["party"] == "Ledger Brand"
    assert (movement["quantity_in"], movement["balance"]) == (40, 40)


def test_range_opens_on_prior_stock_and_balances_run_through_it(db):
    _, warehouse, retailer, sku = _seed(db)
    sale = Order(from_entity_type="WAREHOUSE", from_entity_id=warehouse.id, to_entity_type="RETAILER",
                 to_entity_id=retailer.id, status=OrderStatus.DELIVERED)
    db.add(sale)
    db.commit()
    _move(db, sku, warehouse, TransactionType.IN, 100, datetime(2026, 9, 1, tzinfo=timezone.utc))
    _move(db, sku, warehouse, TransactionType.OUT, 30, datetime(2026, 10, 2, tzinfo=timezone.utc), sale.id)
    _move(db, sku, warehouse, TransactionType.RETURN, 5, datetime(2026, 10, 3, tzinfo=timezone.utc), sale.id)
    _move(db, sku, warehouse, TransactionType.OUT, 10, datetime(2026, 11, 1, tzinfo=timezone.utc))

    ledger = get_stock_ledger(db, from_date="2026-10-01", to_date="2026-10-31")

    summary = ledger["summary"][0]
    assert (summary["opening"], summary["received"], summary["dispatched"], summary["returned"], summary["closing"]) == (100, 0, 30, 5, 75)
    # Newest first, balances carried from the opening stock.
    assert [(m["type_label"], m["balance"]) for m in ledger["movements"]] == [("Return", 75), ("Dispatch", 70)]
    assert ledger["movements"][1]["party"] == "Ledger Retailer"

    dispatches = get_stock_ledger(db, from_date="2026-10-01", to_date="2026-10-31", transaction_type="OUT")
    # A type filter hides rows but never changes the balances shown.
    assert [(m["type_label"], m["balance"]) for m in dispatches["movements"]] == [("Dispatch", 70)]


def test_export_has_summary_ledger_and_filters_sheets(client, db):
    _, warehouse, _, sku = _seed(db)
    _move(db, sku, warehouse, TransactionType.IN, 12, datetime(2026, 10, 5, tzinfo=timezone.utc))
    user = User(email="ledger@admin.test", password_hash=get_password_hash("password"), role=EmployeeRole.ADMIN)
    db.add(user)
    db.commit()
    headers = {"Authorization": f"Bearer {create_access_token({'user_id': user.id, 'role': 'ADMIN'})}"}

    response = client.get(
        f"/admin/inventory/ledger/export?from_date=2026-10-01&to_date=2026-10-31&warehouse_id={warehouse.id}",
        headers=headers,
    )

    assert response.status_code == 200
    workbook = openpyxl.load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Summary", "Ledger", "Filters"]
    assert [c.value for c in workbook["Summary"][2]][-5:] == [0, 12, 0, 0, 12]
    assert workbook["Ledger"]["O2"].value == 12
    assert ["Warehouse", "Ledger WH"] in [[c.value for c in row] for row in workbook["Filters"].iter_rows()]
