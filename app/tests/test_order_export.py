"""Tests for the Orders XLSX export: the invoice number column and the column
autofilter, plus the header/number-format pairing that a column insertion breaks.
"""
import itertools

from openpyxl import load_workbook

from app.models import Brand, Employee, Inventory, Order, OrderItem, SKU, SKUBatch, Retailer, User, Warehouse
from app.models.enums import EmployeeRole, OrderStatus
from app.schemas.order import OrderCreate, OrderItemCreate, StatusUpdate
from app.services.admin import export_orders_excel, get_order_export_rows
from app.services.order import create_outgoing_order, update_order_status

_seed_counter = itertools.count()


def _seed_world(db):
    n = next(_seed_counter)
    brand = Brand(name=f"Export Brand {n}")
    warehouse = Warehouse(name=f"Export WH {n}", location="Delhi", state="Delhi")
    retailer = Retailer(name=f"Export Retailer {n}", state="Delhi")
    db.add_all([brand, warehouse, retailer])
    db.commit()

    sku = SKU(name=f"Export SKU {n}", brand_id=brand.id, hsn_code="12345678",
              mrp=150, sgst_percent=9, cgst_percent=9, igst_percent=18)
    db.add(sku)
    db.commit()

    db.add(SKUBatch(sku_id=sku.id, warehouse_id=warehouse.id, quantity_received=50, remaining_quantity=50))
    db.add(Inventory(sku_id=sku.id, warehouse_id=warehouse.id, total_quantity=50))
    driver = Employee(name=f"Export Driver {n}", email=f"exp-driver-{n}@ascend.com", role=EmployeeRole.DRIVER)
    user = User(email=f"exp-admin-{n}@ascend.com", password_hash="x", role=EmployeeRole.ADMIN)
    db.add_all([driver, user])
    db.commit()
    return warehouse, retailer, sku, driver, user


def _place_order(db, warehouse, retailer, sku, user, quantity):
    return create_outgoing_order(
        db,
        OrderCreate(
            retailer_id=retailer.id,
            warehouse_id=warehouse.id,
            items=[OrderItemCreate(sku_id=sku.id, quantity=quantity, unit_price=100, discount_amount=0)],
        ),
        user,
    )


def _deliver(db, order, driver, user):
    """Dispatch issues the invoice at OUT_FOR_DELIVERY, so a DELIVERED order has one."""
    order.delivery_driver_id = driver.id
    db.commit()
    update_order_status(db, order.id, StatusUpdate(status="READY_TO_SHIP"), user)
    update_order_status(db, order.id, StatusUpdate(status="OUT_FOR_DELIVERY"), user)
    update_order_status(db, order.id, StatusUpdate(status="DELIVERED"), user)
    db.refresh(order)


def test_export_rows_carry_invoice_number_and_blank_when_uninvoiced(db):
    warehouse, retailer, sku, driver, user = _seed_world(db)
    invoiced = _place_order(db, warehouse, retailer, sku, user, 5)
    _deliver(db, invoiced, driver, user)
    pending = _place_order(db, warehouse, retailer, sku, user, 2)

    rows = get_order_export_rows(db)
    by_order = {row["Order ID"]: row for row in rows}

    assert by_order[invoiced.id]["Invoice Number"] == invoiced.invoice.invoice_number
    assert by_order[pending.id]["Invoice Number"] is None


def test_invoice_join_does_not_duplicate_order_lines(db):
    warehouse, retailer, sku, driver, user = _seed_world(db)
    order = _place_order(db, warehouse, retailer, sku, user, 5)
    _deliver(db, order, driver, user)

    rows = [row for row in get_order_export_rows(db) if row["Order ID"] == order.id]
    assert len(rows) == len(order.items)


def test_workbook_has_autofilter_and_formats_aligned_to_headers(db):
    warehouse, retailer, sku, driver, user = _seed_world(db)
    order = _place_order(db, warehouse, retailer, sku, user, 5)
    _deliver(db, order, driver, user)

    sheet = load_workbook(export_orders_excel(db)).active
    headers = [cell.value for cell in sheet[1]]

    assert headers == [
        "Created At",
        "Order Date",
        "Order ID",
        "Invoice Number",
        "Customer Name ( Retailer Name)",
        "Retailer ID",
        "SKU",
        "SKU Quantity",
        "Salesman Name",
        "Salesman Number",
        "MRP",
        "Discount %",
        "Amount",
        "Rate",
    ]
    # Filter dropdowns must span every column, header row included.
    assert sheet.auto_filter.ref == sheet.dimensions
    assert sheet.freeze_panes == "A2"

    # Number formats are addressed by header name, so inserting a column cannot
    # silently shift a money format onto the wrong column.
    col = {header: idx for idx, header in enumerate(headers, start=1)}
    assert sheet.cell(row=2, column=col["Created At"]).number_format == "yyyy-mm-dd hh:mm"
    assert sheet.cell(row=2, column=col["Order Date"]).number_format == "yyyy-mm-dd"
    for header in ("SKU Quantity", "MRP", "Discount %", "Amount", "Rate"):
        assert sheet.cell(row=2, column=col[header]).number_format == "0.00"
    assert sheet.cell(row=2, column=col["Invoice Number"]).value == order.invoice.invoice_number
