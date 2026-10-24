"""Stock ledger: every inventory movement with a running balance, plus an opening / in / out /
returns / closing summary per SKU and warehouse for a date range.

inventory_transactions is complete: every write to Inventory.total_quantity (receipt, dispatch,
cancel-after-dispatch, credit-note restock) adds a row there, so summing the signed rows
reproduces stock at any point in time.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, case, false, func, true
from sqlalchemy.orm import Session

from app.models import (
    SKU,
    Brand,
    CreditNote,
    InventoryTransaction,
    Invoice,
    Order,
    Retailer,
    SKUBatch,
    Warehouse,
)
from app.models.enums import TransactionType
from app.services.admin import _naive_datetime, _parse_date

TYPE_LABELS = {"IN": "Receipt", "OUT": "Dispatch", "RETURN": "Return"}

# IN and RETURN add stock, OUT removes it.
_signed_quantity = case(
    (InventoryTransaction.transaction_type == TransactionType.OUT, -InventoryTransaction.quantity),
    else_=InventoryTransaction.quantity,
)


def _type_value(value):
    return getattr(value, "value", value)


def _scope(query, warehouse_id=None, brand_id=None, sku_id=None):
    if warehouse_id:
        query = query.filter(InventoryTransaction.warehouse_id == warehouse_id)
    if sku_id:
        query = query.filter(InventoryTransaction.sku_id == sku_id)
    if brand_id:
        query = query.filter(SKU.brand_id == brand_id)
    return query


def _summary(db: Session, from_dt, to_dt, warehouse_id, brand_id, sku_id):
    in_range = InventoryTransaction.created_at >= from_dt if from_dt is not None else true()
    before = InventoryTransaction.created_at < from_dt if from_dt is not None else false()
    kind = InventoryTransaction.transaction_type
    query = (
        db.query(
            InventoryTransaction.sku_id,
            InventoryTransaction.warehouse_id,
            func.coalesce(func.sum(case((before, _signed_quantity), else_=0)), 0),
            func.coalesce(func.sum(case((and_(in_range, kind == TransactionType.IN), InventoryTransaction.quantity), else_=0)), 0),
            func.coalesce(func.sum(case((and_(in_range, kind == TransactionType.OUT), InventoryTransaction.quantity), else_=0)), 0),
            func.coalesce(func.sum(case((and_(in_range, kind == TransactionType.RETURN), InventoryTransaction.quantity), else_=0)), 0),
        )
        .join(SKU, SKU.id == InventoryTransaction.sku_id)
        .group_by(InventoryTransaction.sku_id, InventoryTransaction.warehouse_id)
    )
    if to_dt is not None:
        query = query.filter(InventoryTransaction.created_at < to_dt)
    query = _scope(query, warehouse_id, brand_id, sku_id)
    rows = {}
    for sku_id_, warehouse_id_, opening, received, dispatched, returned in query.all():
        opening, received, dispatched, returned = int(opening), int(received), int(dispatched), int(returned)
        if not (opening or received or dispatched or returned):
            continue
        rows[(sku_id_, warehouse_id_)] = {
            "sku_id": sku_id_,
            "warehouse_id": warehouse_id_,
            "opening": opening,
            "received": received,
            "dispatched": dispatched,
            "returned": returned,
            "closing": opening + received - dispatched + returned,
        }
    return rows


def get_stock_ledger(
    db: Session,
    from_date: str | None = None,
    to_date: str | None = None,
    warehouse_id: int | None = None,
    brand_id: int | None = None,
    sku_id: int | None = None,
    transaction_type: str | None = None,
):
    from_dt = _parse_date(from_date, False)
    to_dt = _parse_date(to_date, True)
    summary = _summary(db, from_dt, to_dt, warehouse_id, brand_id, sku_id)

    query = (
        db.query(InventoryTransaction, SKU, Warehouse, SKUBatch, Order, CreditNote, Invoice)
        .join(SKU, SKU.id == InventoryTransaction.sku_id)
        .join(Warehouse, Warehouse.id == InventoryTransaction.warehouse_id)
        .outerjoin(SKUBatch, SKUBatch.id == InventoryTransaction.batch_id)
        .outerjoin(Order, Order.id == InventoryTransaction.order_id)
        .outerjoin(CreditNote, CreditNote.id == InventoryTransaction.credit_note_id)
        .outerjoin(Invoice, Invoice.order_id == InventoryTransaction.order_id)
    )
    if from_dt is not None:
        query = query.filter(InventoryTransaction.created_at >= from_dt)
    if to_dt is not None:
        query = query.filter(InventoryTransaction.created_at < to_dt)
    query = _scope(query, warehouse_id, brand_id, sku_id)
    results = query.order_by(InventoryTransaction.created_at.asc(), InventoryTransaction.id.asc()).all()

    skus = {s.id: s for s in db.query(SKU).filter(SKU.id.in_({key[0] for key in summary}))} if summary else {}
    brand_ids = {sku.brand_id for sku in skus.values()} | {sku.brand_id for _, sku, *_ in results} | {
        order.from_entity_id for *_, order, _, _ in results if order is not None and order.from_entity_type == "BRAND"
    }
    brands = {b.id: b.name for b in db.query(Brand).filter(Brand.id.in_(brand_ids))} if brand_ids else {}
    retailer_ids = {
        order.to_entity_id for *_, order, _, _ in results if order is not None and order.to_entity_type == "RETAILER"
    }
    retailers = {r.id: r.name for r in db.query(Retailer).filter(Retailer.id.in_(retailer_ids))} if retailer_ids else {}

    # Running balance per SKU and warehouse, starting from the opening stock of the range. Type
    # filtering happens after, so a filtered view still shows true balances.
    balances = {key: row["opening"] for key, row in summary.items()}
    movements = []
    for txn, sku, warehouse, batch, order, credit_note, invoice in results:
        kind = _type_value(txn.transaction_type)
        key = (txn.sku_id, txn.warehouse_id)
        signed = -txn.quantity if kind == "OUT" else txn.quantity
        balances[key] = balances.get(key, 0) + signed
        if transaction_type and kind != transaction_type.upper():
            continue
        party = None
        if order is not None:
            if order.from_entity_type == "BRAND":
                party = brands.get(order.from_entity_id)
            elif order.to_entity_type == "RETAILER":
                party = retailers.get(order.to_entity_id)
        movements.append(
            {
                "id": txn.id,
                "created_at": txn.created_at,
                "warehouse_id": txn.warehouse_id,
                "warehouse_name": warehouse.name,
                "sku_id": sku.id,
                "sku_code": sku.code,
                "sku_name": sku.name,
                "brand_name": brands.get(sku.brand_id),
                "batch_id": txn.batch_id,
                "expiry_date": batch.expiry_date if batch else None,
                "transaction_type": kind,
                "type_label": TYPE_LABELS.get(kind, kind),
                "order_id": txn.order_id,
                "party": party,
                "invoice_number": invoice.invoice_number if invoice else None,
                "credit_note_number": credit_note.credit_note_number if credit_note else None,
                "quantity_in": signed if signed > 0 else 0,
                "quantity_out": -signed if signed < 0 else 0,
                "balance": balances[key],
            }
        )

    warehouses = {w.id: w.name for w in db.query(Warehouse)}
    summary_rows = []
    for (sku_key, warehouse_key), row in summary.items():
        sku = skus.get(sku_key)
        summary_rows.append(
            {
                **row,
                "warehouse_name": warehouses.get(warehouse_key),
                "sku_code": sku.code if sku else None,
                "sku_name": sku.name if sku else None,
                "brand_name": brands.get(sku.brand_id) if sku else None,
            }
        )
    summary_rows.sort(key=lambda r: (r["warehouse_name"] or "", r["brand_name"] or "", r["sku_name"] or ""))
    # Newest first on screen; the export re-sorts oldest first so balances read top to bottom.
    movements.reverse()
    return {"summary": summary_rows, "movements": movements}


def _write_sheet(sheet, headers, rows, formats=None):
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for col_idx, fmt in (formats or {}).items():
        for cell in sheet.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            cell[0].number_format = fmt
    widths = [len(header) for header in headers]
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=len(headers)):
        for idx, cell in enumerate(row):
            text = "" if cell.value is None else str(cell.value)
            widths[idx] = max(widths[idx], len(text))
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 40)


def export_stock_ledger_xlsx(db: Session, filters: dict, filter_labels: dict) -> BytesIO:
    data = get_stock_ledger(db, **filters)
    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    _write_sheet(
        summary_sheet,
        ["Warehouse", "Brand", "SKU Code", "SKU", "Opening", "Received", "Dispatched", "Returned", "Closing"],
        [
            [r["warehouse_name"], r["brand_name"], r["sku_code"], r["sku_name"], r["opening"], r["received"], r["dispatched"], r["returned"], r["closing"]]
            for r in data["summary"]
        ],
    )

    ledger_sheet = workbook.create_sheet("Ledger")
    _write_sheet(
        ledger_sheet,
        ["Date", "Warehouse", "Brand", "SKU Code", "SKU", "Batch", "Expiry", "Type", "Order ID", "Party", "Invoice", "Credit Note", "In", "Out", "Balance"],
        [
            [
                _naive_datetime(m["created_at"]), m["warehouse_name"], m["brand_name"], m["sku_code"], m["sku_name"],
                m["batch_id"], m["expiry_date"], m["type_label"], m["order_id"], m["party"], m["invoice_number"],
                m["credit_note_number"], m["quantity_in"] or None, m["quantity_out"] or None, m["balance"],
            ]
            for m in reversed(data["movements"])
        ],
        formats={1: "yyyy-mm-dd hh:mm", 7: "yyyy-mm-dd"},
    )

    filters_sheet = workbook.create_sheet("Filters")
    _write_sheet(filters_sheet, ["Filter", "Value"], [[label, value] for label, value in filter_labels.items()])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
